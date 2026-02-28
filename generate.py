# 旧版：需求 → 六维测试点（功能/边界/异常/权限/性能/隐藏雷点）JSON + Excel + XMind，另一套结构。
import os
import json
import sys
from collections import defaultdict

from dotenv import load_dotenv
from jsonschema import validate
from openai import OpenAI
from openpyxl import Workbook
from xmind_export_zen import save_to_xmind_zen_by_module_reviewable
load_dotenv()

# 使用阿里通义千问（DashScope）的 API Key
API_KEY = os.getenv("DASHSCOPE_API_KEY")
if not API_KEY:
    raise SystemExit("Missing DASHSCOPE_API_KEY in .env")

# DashScope 的 OpenAI 兼容接口：
# - 中国站常用 base_url: https://dashscope.aliyuncs.com/compatible-mode/v1
# - 国际站示例 base_url: https://dashscope-intl.aliyuncs.com/compatible-mode/v1
BASE_URL = os.getenv(
    "DASHSCOPE_BASE_URL",
    "https://dashscope.aliyuncs.com/compatible-mode/v1",
)

client = OpenAI(
    api_key=API_KEY,
    base_url=BASE_URL,
)

SCHEMA = {
    "type": "object",
    "properties": {
        "功能测试": {"type": "array", "items": {"type": "string"}},
        "边界测试": {"type": "array", "items": {"type": "string"}},
        "异常测试": {"type": "array", "items": {"type": "string"}},
        "权限测试": {"type": "array", "items": {"type": "string"}},
        "性能与容量风险": {"type": "array", "items": {"type": "string"}},
        "隐藏雷点提示": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["功能测试", "边界测试", "异常测试", "权限测试", "性能与容量风险", "隐藏雷点提示"],
    "additionalProperties": False,
}

# 额外业务上下文文件（可选）
CONTEXT_GOODS_PATH = "context_goods.txt"


def read_text(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read().strip()


def llm_generate_test_points(req_text: str) -> dict:
    # 读取可选的商品模块业务上下文
    goods_ctx = ""
    if os.path.exists(CONTEXT_GOODS_PATH):
        goods_ctx = read_text(CONTEXT_GOODS_PATH)

    prompt_parts = [
        "你是资深软件测试工程师。",
    ]

    if goods_ctx:
        prompt_parts.append("下面是【商品模块业务上下文】（请充分利用其中的信息）：")
        prompt_parts.append(goods_ctx)

    prompt_parts.append("下面是【本次需求文档正文】：")
    prompt_parts.append(req_text)

    prompt_parts.append(
        """
请输出“测试点清单”，必须严格符合 JSON 格式与字段，不要输出任何多余文字。

字段要求：
- 功能测试：功能正向/流程
- 边界测试：数量/长度/范围/分页/批量等边界
- 异常测试：参数非法/缺失/并发/网络/重复提交等
- 权限测试：不同角色、越权、数据范围
- 性能与容量风险：大数据量、导出、批量操作等
- 隐藏雷点提示：容易漏测/上线事故高发点（请更狠一点）
""".strip()
    )

    prompt = "\n\n".join(prompt_parts)

    resp = client.chat.completions.create(
        model="qwen-plus",
        messages=[
            {"role": "system", "content": "You output ONLY valid JSON."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.2,
    )
    text = resp.choices[0].message.content
    data = json.loads(text)
    validate(instance=data, schema=SCHEMA)
    return data


def save_to_excel(data: dict, out_path: str) -> None:
    """将测试点 JSON 保存为 Excel：分类 + 序号 + 测试点描述"""
    wb = Workbook()
    ws = wb.active
    ws.title = "测试点"

    # 表头
    ws.append(["分类", "序号", "测试点"])

    for category, tests in data.items():
        if not isinstance(tests, list):
            continue
        for idx, item in enumerate(tests, start=1):
            ws.append([category, idx, str(item)])

    wb.save(out_path)


# ----------------------------
# XMind：按“模块/页面”评审结构输出
# ----------------------------
def _infer_module(item: str, module_rules: list[tuple[str, list[str]]]) -> str:
    """
    根据关键词规则把测试点归到某个模块/页面
    module_rules: [("处理完成页", ["处理完成", "已处理", "完成列表"]), ...]
    """
    s = str(item)
    for module_name, keywords in module_rules:
        for kw in keywords:
            if kw and kw in s:
                return module_name
    return "未归类"


def save_to_xmind_by_module_reviewable(
    data: dict,
    out_path: str,
    root_title: str = "测试点清单",
    module_rules: list[tuple[str, list[str]]] | None = None,
    add_placeholders: bool = True,
    print_uncategorized: bool = True,
) -> None:
    """
    评审友好 XMind（按模块/页面）：
    根 -> 模块/页面 -> 6维度 -> 测试点(编号) -> 占位(可选)
    """

    category_order = [
        "功能测试",
        "边界测试",
        "异常测试",
        "权限测试",
        "性能与容量风险",
        "隐藏雷点提示",
    ]
    cat_code = {
        "功能测试": "FUNC",
        "边界测试": "BND",
        "异常测试": "ERR",
        "权限测试": "AUTH",
        "性能与容量风险": "PERF",
        "隐藏雷点提示": "RISK",
    }

    # 默认规则：你可以按你系统页面补充/改名
    if module_rules is None:
        module_rules = [
            ("列表页", ["列表", "表格", "筛选", "排序", "分页", "查询", "搜索"]),
            ("详情页/弹窗", ["详情", "弹窗", "对话框", "抽屉", "确认", "提示"]),
            ("导出/下载", ["导出", "下载", "文件", "任务", "数据导出"]),
            ("批量操作", ["批量", "多选", "勾选", "全选", "反选"]),
            ("设置/配置", ["设置", "配置", "规则", "开关"]),
            ("接口/后端", ["接口", "API", "请求", "入参", "返回", "字段"]),
        ]

    # 整理结构：module -> category -> [items]
    buckets: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for category in category_order:
        tests = data.get(category, [])
        if not isinstance(tests, list):
            continue
        for item in tests:
            module = _infer_module(str(item), module_rules)
            buckets[module][category].append(str(item))

    # 打印未归类统计，方便你迭代 module_rules
    if print_uncategorized:
        uncat_count = sum(len(v) for v in buckets.get("未归类", {}).values())
        print(f"[XMind] 未归类测试点数量: {uncat_count}")

    # 载入/创建 xmind
    workbook = xmind.load(out_path)
    sheet = workbook.getPrimarySheet()
    sheet.setTitle("测试点评审-按模块")

    root = sheet.getRootTopic()
    root.setTitle(root_title)

    # 模块排序：未归类放最后，其它按字典序
    modules = sorted(buckets.keys(), key=lambda x: (x == "未归类", x))

    for module in modules:
        mod_topic = root.addSubTopic()
        mod_topic.setTitle(module)

        # 模块内编号从 1 开始（更符合“页面内评审”）
        module_seq = 0

        for category in category_order:
            items = buckets[module].get(category, [])
            if not items:
                continue

            cat_topic = mod_topic.addSubTopic()
            cat_topic.setTitle(category)

            code = cat_code.get(category, "GEN")

            for item in items:
                module_seq += 1
                tp_id = f"TP-{code}-{module_seq:03d}"

                title = f"{tp_id}  {item}"
                if category == "隐藏雷点提示":
                    title = f"⚠ {title}"

                tp_topic = cat_topic.addSubTopic()
                tp_topic.setTitle(title)

                if add_placeholders:
                    tp_topic.addSubTopic().setTitle("优先级：P0/P1/P2（待定）")
                    tp_topic.addSubTopic().setTitle("备注：口径/数据/前置（待补）")
                    tp_topic.addSubTopic().setTitle("关联：页面/接口/字段（待补）")

    xmind.save(workbook, out_path)


def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python generate.py inputs/req.md")

    in_path = sys.argv[1]
    req_text = read_text(in_path)

    result = llm_generate_test_points(req_text)

    os.makedirs("outputs", exist_ok=True)

    # 保存 JSON
    json_path = os.path.join("outputs", "test_points.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    # 保存 Excel
    excel_path = os.path.join("outputs", "test_points.xlsx")
    save_to_excel(result, excel_path)

    # 保存 XMind（ZEN 格式，按模块/页面，评审结构）
    xmind_path = os.path.join("outputs", "test_points.xmind")
    save_to_xmind_zen_by_module_reviewable(
        result,
        xmind_path,
        root_title=os.path.basename(in_path),
        module_rules=[
            ("重量管理-称重异常-处理完成页", ["处理完成", "已处理", "完成页面", "处理错误", "调整"]),
            ("批量搜索弹窗", ["批量搜索", "输入sku", "输入spu", "确定", "弹窗"]),
            ("导出按钮/导出任务页", ["导出", "导出任务", "下载", "文件", "任务列表"]),
            ("重量修改确认/包装商品提示", ["修改重量", "包装商品", "提示校验", "确认修改", "校验"]),
            ("列表筛选/分页", ["筛选", "分页", "排序", "搜索", "查询", "条件项"]),
            ("权限/数据范围", ["权限", "角色", "越权", "数据范围", "可见"]),
        ],
        add_placeholders=True,
        print_uncategorized=True,
    )
    print(f"OK -> {xmind_path}")



if __name__ == "__main__":
    main()